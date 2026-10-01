"""Builds SIH2026_PS26138_Idea.pptx on the official SIH template.

Keeps the template's title, logo, team oval, footer and section pointers (wording unchanged, restyled small),
and fills the content area with native shapes and native tables fed from results.json (python bench.py).
"""
import copy
import io
import json
import os
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.dirname(ROOT)
TEMPLATE = os.path.join(PROJ, "SIH2026-IDEA-Presentation-Format.pptx")
OUT = os.path.join(PROJ, "SIH2026_PS26138_Idea.pptx")
TEAM = "Prometheus 01"

NAVY = RGBColor(0x0F, 0x2B, 0x4C)
BLUE = RGBColor(0x00, 0x70, 0xC0)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
ORANGE = RGBColor(0xD9, 0x82, 0x2B)
TINT = RGBColor(0xEE, 0xF5, 0xFB)
TINT2 = RGBColor(0xFD, 0xF3, 0xE7)
GREY = RGBColor(0x55, 0x65, 0x75)
LINE = RGBColor(0xC9, 0xD6, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FOAM = RGBColor(0x9F, 0xE7, 0xDA)
FONT = "Calibri"

# ---------------------------------------------------------------- numbers
RES = json.load(open(os.path.join(ROOT, "results.json")))


def fmt_s(t):
    return f"{t:.2f} s" if t < 1 else f"{t:.1f} s"


def row_text(r):
    exact = f"Proven best · {fmt_s(r['milp_time_s'])}"
    if r["qi_feasible"]:
        qi = f"+{r['qi_gap_pct']:.1f}% · {fmt_s(r['qi_time_s'])}"
    else:
        qi = "No plan within fleet limits"
    return exact, qi


# ---------------------------------------------------------------- helpers
def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, paras, size=12, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=4, italic=False, shape=None):
    """paras: list of paragraphs; each is a str or a list of (text, {overrides}) runs."""
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        runs = [(p, {})] if isinstance(p, str) else p
        for t, o in runs:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return shape


def boxed_text(slide, x, y, w, h, paras, fill, pad=0.16, line=None, anchor=MSO_ANCHOR.TOP, **kw):
    b = box(slide, x, y, w, h, fill=fill, line=line)
    tf = b.text_frame
    text(slide, 0, 0, 0, 0, paras, anchor=anchor, shape=b, **kw)
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(pad * 0.8)
    return b


def circle_num(slide, x, y, n, fill):
    c = box(slide, x, y, 0.42, 0.42, fill=fill, shape=MSO_SHAPE.OVAL)
    text(slide, 0, 0, 0, 0, [str(n)], size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=c)
    return c


def set_team(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "Your Team Name":
            p = sh.text_frame.paragraphs[0]
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
            p.runs[0].text = TEAM
            p.runs[0].font.size = Pt(11); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = NAVY
            sh.text_frame.margin_left = sh.text_frame.margin_right = 0; sh.text_frame.word_wrap = False


def restyle_pointers(slide):
    """Keep the template's pointer text unchanged but move it into one small grey line under the title."""
    title = slide.shapes.title
    for sh in list(slide.shapes):
        if sh is title or not sh.has_text_frame or sh.is_placeholder:
            continue
        t = sh.text_frame.text
        if any(k in t for k in ["Proposed Solution", "Technologies to be used", "Analysis of the feasibility",
                                "Potential impact", "Details / Links"]):
            items = [p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip()]
            sh._element.getparent().remove(sh._element)
            text(slide, 1.75, 1.2, 10.3, 0.3, [" · ".join(items)], size=9, color=GREY, italic=True)
            return


def heading(slide, s, color=NAVY, y=1.6, size=20):
    text(slide, 0.5, y, 12.3, 0.5, [s], size=size, bold=True, color=color)


# ---------------------------------------------------------------- build
prs = Presentation(TEMPLATE)

# drop the "Important instructions" slide (template says it may be deleted)
sldIdLst = prs.slides._sldIdLst
last = sldIdLst[-1]
prs.part.drop_rel(last.rId)
sldIdLst.remove(last)

s1, s2, s3, s4, s5, s6 = prs.slides

# ---------- slide 1: title page
for sh in s1.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "TITLE PAGE":
        tf = sh.text_frame
        p = [q for q in tf.paragraphs if q.runs and "TITLE PAGE" in q.text][0]
        for q in tf.paragraphs:
            if q._p is not p._p:
                q._p.getparent().remove(q._p)
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].text = "SagarMitra: exact where it wins, quantum-inspired where it has to be"
        p.runs[0].font.size = Pt(18); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = TEAL
        sh.top = sh.top + Inches(0.3)
    if sh.has_text_frame and "Problem Statement ID" in sh.text_frame.text:
        values = {
            "Problem Statement ID": "26138",
            "Problem Statement Title": "Quantum-Inspired Fuel Consumption Prediction and Green Fleet Optimization",
            "Theme": "Clean & Green Technology",
            "PS Category": "Software",
            "Team ID": "151347",
            "Team Name": TEAM,
        }
        for p in sh.text_frame.paragraphs:
            t = p.text
            for key, val in values.items():
                if t.startswith(key):
                    label = {"PS Category": "PS Category: ", "Team Name": "Team Name (Registered on portal): "}.get(key, key + ": ")
                    base = p.runs[0]
                    for r in p.runs[1:]:
                        r._r.getparent().remove(r._r)
                    base.text = label
                    base.font.bold = True; base.font.size = Pt(15); base.font.color.rgb = NAVY
                    v = copy.deepcopy(base._r); base._r.addnext(v)
                    vr = p.runs[1]; vr.text = val; vr.font.bold = False
                    break

# Prometheus logo from the team's earlier deck, if present
old = os.path.join(PROJ, "SIH2026_PS26034_Idea.pptx")
if os.path.exists(old):
    with zipfile.ZipFile(old) as z:
        if "ppt/media/image3.png" in z.namelist():
            s1.shapes.add_picture(io.BytesIO(z.read("ppt/media/image3.png")), Inches(0.35), Inches(0.25), Inches(0.9), Inches(0.9))

exec(open(os.path.join(ROOT, "deck_slides_2_4.py"), encoding="utf-8").read())

# ---------- slide 5: impact
set_team(s5); restyle_pointers(s5)
text(s5, 0.5, 1.6, 6, 0.35, ["Who uses it"], size=14, bold=True, color=TEAL)
users = [
    ("Shipping lines and PSU fleets", "Cheaper plans that already meet the IMO carbon rating."),
    ("Port authorities", "Where shore power and green-fuel bunkering would serve the most ships."),
    ("Ministry of Ports, Shipping and Waterways", "Test fleet-transition plans for Maritime India Vision 2030 before committing money."),
]
for i, (t, d) in enumerate(users):
    x = 0.5 + i * 4.175
    box(s5, x, 1.98, 3.95, 1.3, fill=TINT)
    text(s5, x + 0.2, 2.1, 3.55, 0.45, [t], size=14, bold=True, color=NAVY)
    text(s5, x + 0.2, 2.58, 3.55, 0.65, [d], size=11.5, color=GREY)

text(s5, 0.5, 3.48, 6, 0.35, ["What changes"], size=14, bold=True, color=TEAL)
tiles = [
    ("Economic", ORANGE, "Lower fuel bills. Sailing slower where the schedule allows cuts fuel per nautical mile sharply."),
    ("Environmental", TEAL, "Lifecycle emissions, so a plan cannot look green by moving emissions upstream."),
    ("Regulatory", BLUE, "Every plan meets the IMO carbon intensity rating before a planner sees it."),
    ("Research", NAVY, "An open, repeatable test of quantum-inspired against exact methods on an Indian network."),
]
for i, (t, col, d) in enumerate(tiles):
    x = 0.5 + i * 3.1
    box(s5, x, 3.85, 2.9, 1.8, fill=WHITE, line=LINE)
    text(s5, x + 0.2, 3.98, 2.5, 0.4, [t], size=15, bold=True, color=col)
    text(s5, x + 0.2, 4.42, 2.5, 1.15, [d], size=11.5, color=GREY)
boxed_text(s5, 0.5, 5.85, 12.3, 0.55,
           [[("Future: ", {"bold": True, "color": FOAM}), ("the same engine plans truck and rail fleets, and any routing problem with a mix of linear and nonlinear costs.", {})]],
           fill=NAVY, size=12.5, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, pad=0.2)

# ---------- slide 6: references
set_team(s6); restyle_pointers(s6)
cols = [
    ("Methods", [
        "Han & Kim (2002). Quantum-inspired evolutionary algorithm for a class of combinatorial optimization. IEEE Trans. Evolutionary Computation.",
        "Sun, Feng & Xu (2004). Particle swarm optimization with particles having quantum behavior. IEEE CEC.",
        "Huangfu & Hall (2018). Parallelizing the dual revised simplex method. Mathematical Programming Computation (HiGHS).",
        "Deb et al. (2002). A fast and elitist multiobjective genetic algorithm: NSGA-II. IEEE Trans. Evolutionary Computation.",
        "Psaraftis & Kontovas (2013). Speed models for energy-efficient maritime transportation. Transportation Research Part C.",
    ]),
    ("Standards, policy and data", [
        "IMO MEPC.376(80), 2023: Guidelines on life cycle GHG intensity of marine fuels (LCA guidelines).",
        "IMO 2023 Strategy on reduction of GHG emissions from ships; CII rating (MARPOL Annex VI).",
        "Ministry of Ports, Shipping and Waterways: Maritime India Vision 2030; Harit Sagar Green Port Guidelines (2023).",
        "EU-MRV / THETIS ship emissions data: mrv.emsa.europa.eu",
        "Open AIS vessel tracks; Copernicus ERA5 weather reanalysis.",
    ]),
]
for i, (h, items) in enumerate(cols):
    x = 0.5 + i * 6.25
    box(s6, x, 1.65, 6.05, 3.85, fill=TINT)
    text(s6, x + 0.25, 1.82, 5.6, 0.4, [h], size=15, bold=True, color=TEAL)
    text(s6, x + 0.25, 2.3, 5.6, 4.0, ["•  " + t for t in items], size=11.5, color=NAVY, space_after=10)

boxed_text(s6, 0.5, 5.72, 12.3, 0.62,
           [[("Our prototype: ", {"bold": True, "color": FOAM}),
             ("fleet model, HiGHS exact solver, QIEA + QPSO and the benchmark script, in ps26138-sagarmitra (fleet.py, solvers.py, bench.py).", {})]],
           fill=NAVY, size=12.5, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, pad=0.2)

prs.save(OUT)
print("saved", OUT)
